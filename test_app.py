import io
import json
import unittest
from pathlib import Path

import docx
from docx import Document
from docx.enum.text import WD_COLOR_INDEX

import app


class TestLegalTemplateGeneratorOllama(unittest.TestCase):

    def setUp(self):
        self.config = app.load_config()

    def test_config_structure(self):
        self.assertIn("system_prompt", self.config)
        self.assertIn("variables", self.config)
        self.assertIn("base_url", self.config)
        self.assertIn("selected_model", self.config)
        self.assertIn("CRITICAL RULE", self.config["system_prompt"])
        var_names = app.get_variable_names(self.config["variables"])
        self.assertIn("pt_agreement_number", var_names)
        self.assertIn("pt_first_name", var_names)
        self.assertIn("pt_total_debt", var_names)

    def test_docx_text_extraction(self):
        doc = Document()
        doc.add_paragraph("Договір позики № 987654.")
        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Позичальник:"
        table.cell(0, 1).text = "Шевченко Тарас"

        extracted = app.extract_text_from_docx(doc)
        self.assertIn("987654", extracted)
        self.assertIn("Позичальник:", extracted)
        self.assertIn("Шевченко Тарас", extracted)

    def test_docx_replacement_and_highlight(self):
        doc = Document()
        p = doc.add_paragraph()
        p.add_run("Договір позики № ")
        p.add_run("4820")
        p.add_run("-2024")

        table = doc.add_table(rows=1, cols=2)
        table.cell(0, 0).text = "Сума:"
        table.cell(0, 1).text = "52300 грн."

        replacements = [
            ("4820-2024", "pt_agreement_number"),
            ("52300", "pt_total_debt")
        ]

        mod_doc, count = app.generate_template_docx(doc, replacements)
        self.assertEqual(count, 2)

        # Verify paragraph replacement & highlight
        p_reloaded = mod_doc.paragraphs[0]
        self.assertIn("$$$pt_agreement_number$$$", p_reloaded.text)
        highlighted_runs = [r for r in p_reloaded.runs if r.font.highlight_color == WD_COLOR_INDEX.YELLOW]
        self.assertEqual(len(highlighted_runs), 1)
        self.assertEqual(highlighted_runs[0].text, "$$$pt_agreement_number$$$")

        # Verify table cell replacement & highlight
        cell_p = mod_doc.tables[0].cell(0, 1).paragraphs[0]
        self.assertIn("$$$pt_total_debt$$$", cell_p.text)
        cell_highlighted = [r for r in cell_p.runs if r.font.highlight_color == WD_COLOR_INDEX.YELLOW]
        self.assertEqual(len(cell_highlighted), 1)
        self.assertEqual(cell_highlighted[0].text, "$$$pt_total_debt$$$")

    def test_in_memory_buffer_save(self):
        doc = Document()
        doc.add_paragraph("Тестовий документ")
        buf = io.BytesIO()
        doc.save(buf)
        buf.seek(0)
        self.assertGreater(len(buf.getvalue()), 0)
        reloaded = Document(buf)
        self.assertEqual(reloaded.paragraphs[0].text, "Тестовий документ")

    def test_json_cleaning(self):
        raw_response = "```json\n{\n  \"4820-2024\": \"pt_agreement_number\"\n}\n```"
        cleaned = raw_response.strip()
        if cleaned.startswith("```"):
            import re
            cleaned = re.sub(r"^```(?:json)?\s*", "", cleaned)
            cleaned = re.sub(r"\s*```$", "", cleaned)
        data = json.loads(cleaned)
        self.assertEqual(data.get("4820-2024"), "pt_agreement_number")

    def test_fetch_available_models_graceful_on_unreachable_endpoint(self):
        # Must return empty list gracefully and not raise uncaught exception
        models = app.fetch_available_models("http://127.0.0.1:9999/v1")
        self.assertIsInstance(models, list)
        self.assertEqual(len(models), 0)


if __name__ == "__main__":
    unittest.main()
