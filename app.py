import copy
import io
import json
import os
import re
from pathlib import Path
from typing import Any, Dict, List, Optional, Tuple

import docx
from docx import Document
from docx.enum.text import WD_COLOR_INDEX
import google.generativeai as genai
from openai import OpenAI
import pandas as pd
import streamlit as st

# ==========================================
# CONFIGURATION & CONSTANTS
# ==========================================
CONFIG_FILE = Path(__file__).parent / "config.json"

DEFAULT_PROVIDER = "Local LLM (Ollama)"
DEFAULT_LOCAL_BASE_URL = "http://10.0.1.17:4333/v1"
DEFAULT_LOCAL_MODEL = "qwen3.8:27b"
DEFAULT_GEMINI_MODEL = "models/gemini-3.5-flash-lite"

DEFAULT_SYSTEM_PROMPT = (
    "Ти — експерт-юрист. Твоє завдання — аналізувати текст і знаходити значення, "
    "що відповідають глобальним змінним. ІГНОРУЙ константи: не чіпай реквізити "
    "відправника (якщо вони незмінні), індекси, загальні номери законів. "
    "CRITICAL RULE: The keys in your JSON response MUST be exact, "
    "character-by-character copies of the text found in the document "
    "(including spaces and punctuation). Do not reformat dates or numbers. "
    "Повертай ТІЛЬКИ валідний JSON: ключ — знайдений шматок тексту, значення — назва змінної. "
    "CRITICAL: You must return ONLY raw, valid JSON. Do not include any markdown formatting, "
    "do not use ```json blocks, and do not add any explanations, greetings, or conversational text "
    "before or after the JSON. Just the raw JSON object."
)

DEFAULT_VARIABLES = {
    "pt_address": "Адреса реєстрації або фактичного проживання",
    "pt_agreement_number": "Номер кредитного або іншого договору",
    "pt_capital": "Сума основного боргу (тіло кредиту)",
    "pt_first_name": "Ім'я сторони",
    "pt_last_name": "Прізвище сторони",
    "pt_middle_name": "По батькові сторони",
    "pt_pin": "РНОКПП (ІПН) або паспортні дані",
    "pt_total_debt": "Загальна сума заборгованості",
    "pt_court_fee": "Сума судового збору",
    "pt_phone": "Номер телефону"
}

DEFAULT_CONFIG = {
    "provider": DEFAULT_PROVIDER,
    "gemini_api_key": "",
    "base_url": DEFAULT_LOCAL_BASE_URL,
    "selected_model": DEFAULT_LOCAL_MODEL,
    "system_prompt": DEFAULT_SYSTEM_PROMPT,
    "variables": DEFAULT_VARIABLES
}


def get_variable_names(var_dict: Dict[str, Any]) -> List[str]:
    """Returns the list of valid variable identifiers for the SelectboxColumn."""
    names = []
    for k, v in var_dict.items():
        if isinstance(v, str) and (v.startswith("pt_") or "_" in v):
            names.append(v)
        elif isinstance(k, str) and (k.startswith("pt_") or "_" in k):
            names.append(k)
        else:
            names.append(str(k))
    return sorted(list(set(names))) if names else list(var_dict.keys())


# ==========================================
# CONFIGURATION MANAGEMENT
# ==========================================
def load_config() -> Dict[str, Any]:
    """Loads config.json or creates default if not present."""
    if not CONFIG_FILE.exists():
        save_config(DEFAULT_CONFIG)
        return copy.deepcopy(DEFAULT_CONFIG)
    try:
        with open(CONFIG_FILE, "r", encoding="utf-8") as f:
            data = json.load(f)
            if not isinstance(data, dict):
                data = copy.deepcopy(DEFAULT_CONFIG)

            if "provider" not in data:
                data["provider"] = "Google Gemini" if data.get("gemini_api_key") else DEFAULT_PROVIDER
            if "gemini_api_key" not in data:
                data["gemini_api_key"] = ""
            if "base_url" not in data:
                data["base_url"] = DEFAULT_LOCAL_BASE_URL
            if "selected_model" not in data:
                data["selected_model"] = DEFAULT_GEMINI_MODEL if data.get("provider") == "Google Gemini" else DEFAULT_LOCAL_MODEL

            if "system_prompt" not in data:
                data["system_prompt"] = DEFAULT_SYSTEM_PROMPT
            else:
                directive = (
                    "CRITICAL: You must return ONLY raw, valid JSON. Do not include any markdown formatting, "
                    "do not use ```json blocks, and do not add any explanations, greetings, or conversational text "
                    "before or after the JSON. Just the raw JSON object."
                )
                if directive not in data["system_prompt"]:
                    data["system_prompt"] = data["system_prompt"].strip() + " " + directive

            if "variables" not in data or not isinstance(data["variables"], dict):
                data["variables"] = DEFAULT_VARIABLES

            return data
    except Exception as e:
        print(f"Помилка завантаження config.json: {e}")
        return copy.deepcopy(DEFAULT_CONFIG)


def save_config(config_dict: Dict[str, Any]) -> bool:
    """Saves dictionary to config.json."""
    try:
        with open(CONFIG_FILE, "w", encoding="utf-8") as f:
            json.dump(config_dict, f, ensure_ascii=False, indent=2)
        return True
    except Exception as e:
        print(f"Помилка збереження config.json: {e}")
        return False


# ==========================================
# DYNAMIC MODEL DISCOVERY
# ==========================================
def fetch_gemini_models(api_key: str) -> List[str]:
    """Dynamically discovers all models supporting generateContent from Gemini API."""
    if not api_key or not api_key.strip():
        return []
    try:
        genai.configure(api_key=api_key.strip())
        supported = []
        for m in genai.list_models():
            methods = getattr(m, "supported_generation_methods", [])
            if "generateContent" in methods:
                clean_name = m.name
                if not any(skip in clean_name.lower() for skip in ["tts", "image", "clip", "robotics"]):
                    supported.append(clean_name)
        return supported
    except Exception as e:
        print(f"Помилка отримання моделей Gemini: {e}")
        return []


def fetch_local_models(base_url: str) -> List[str]:
    """Dynamically discovers all models from Local OpenAI-compatible Ollama endpoint."""
    if not base_url or not base_url.strip():
        return []
    try:
        client = OpenAI(
            base_url=base_url.strip(),
            api_key="ollama",
            timeout=5.0
        )
        models_page = client.models.list()
        model_ids = [m.id for m in models_page if hasattr(m, "id")]
        return model_ids
    except Exception as e:
        print(f"Не вдалося отримати моделі з Local LLM ({base_url}): {e}")
        return []


# Backwards compatibility alias
fetch_available_models = fetch_local_models


# ==========================================
# DOCX EXTRACTION & REPLACEMENT ENGINE
# ==========================================
def extract_text_from_docx(doc: Document) -> str:
    """Extracts text from both paragraphs and table cells preserving structure."""
    parts = []

    # 1. Main body paragraphs
    for p in doc.paragraphs:
        t = p.text.strip()
        if t:
            parts.append(t)

    # 2. Table cells
    for table in doc.tables:
        for row in table.rows:
            row_items = []
            for cell in row.cells:
                cell_text = "\n".join(p.text.strip() for p in cell.paragraphs if p.text.strip())
                if cell_text and cell_text not in row_items:
                    row_items.append(cell_text)
            if row_items:
                parts.append(" | ".join(row_items))

    return "\n\n".join(parts)


def replace_in_paragraph(p, replacements: List[Tuple[str, str]]) -> int:
    """
    Robust paragraph-level replacement:
    Replaces original snippets with $$$var$$$, clears existing runs,
    and rebuilds runs with yellow highlighting specifically applied to the variable tags.
    """
    original_text = p.text
    if not original_text:
        return 0

    new_text = original_text
    match_count = 0
    for orig, var in replacements:
        if orig in new_text:
            new_text = new_text.replace(orig, f"$$${var}$$$")
            match_count += 1

    if new_text != original_text:
        font_name = None
        font_size = None
        is_bold = None
        is_italic = None
        color_rgb = None
        if p.runs:
            ref = p.runs[0]
            font_name = ref.font.name
            font_size = ref.font.size
            is_bold = ref.bold
            is_italic = ref.italic
            if ref.font.color and ref.font.color.rgb:
                color_rgb = ref.font.color.rgb

        for r in list(p.runs):
            p._p.remove(r._r)

        tokens = re.split(r"(\$\$\$[a-zA-Z0-9_]+\$\$\$)", new_text)
        for token in tokens:
            if not token:
                continue
            run = p.add_run(token)
            if font_name:
                run.font.name = font_name
            if font_size:
                run.font.size = font_size
            if is_bold is not None:
                run.bold = is_bold
            if is_italic is not None:
                run.italic = is_italic
            if color_rgb is not None:
                run.font.color.rgb = color_rgb

            if re.match(r"^\$\$\$[a-zA-Z0-9_]+\$\$\$$", token):
                run.font.highlight_color = WD_COLOR_INDEX.YELLOW

    return match_count


def generate_template_docx(doc: Document, replacements: List[Tuple[str, str]]) -> Tuple[Document, int]:
    """Applies all replacements with yellow highlights across document body and tables."""
    total_count = 0
    visited_paragraphs = set()

    for p in doc.paragraphs:
        if id(p) not in visited_paragraphs:
            visited_paragraphs.add(id(p))
            total_count += replace_in_paragraph(p, replacements)

    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    if id(p) not in visited_paragraphs:
                        visited_paragraphs.add(id(p))
                        total_count += replace_in_paragraph(p, replacements)

    return doc, total_count


# ==========================================
# TEXT GENERATION LOGIC (HYBRID PROVIDERS)
# ==========================================
def parse_llm_json_response(raw_text: str) -> Dict[str, str]:
    """Cleans markdown blocks from LLM output and parses strict JSON dictionary."""
    text = (raw_text or "").strip()
    if not text:
        raise ValueError("The LLM returned an empty response.")

    # Strip markdown block formatting if present
    if text.startswith("```json"):
        text = text[7:]
    elif text.startswith("```"):
        text = text[3:]
    if text.endswith("```"):
        text = text[:-3]

    text = text.strip()

    try:
        parsed_json = json.loads(text)
    except json.JSONDecodeError as err:
        raise ValueError(f"Не вдалося розпарсити відповідь LLM як JSON: {err}\nОтримана відповідь:\n{raw_text}")

    if not isinstance(parsed_json, dict):
        raise ValueError("Очікувався JSON-об'єкт (dict) формату {'текст': 'назва_змінної'}, але отримано інший тип.")

    return parsed_json


def query_gemini_for_variables(
    api_key: str,
    model_name: str,
    system_prompt: str,
    variables_dict: Dict[str, str],
    document_text: str,
) -> Dict[str, str]:
    """Invokes Google Gemini API for variable extraction."""
    if not api_key.strip():
        raise ValueError("Gemini API Key не вказано. Будь ласка, введіть його у вкладці '⚙️ Settings'.")

    if not model_name.strip():
        raise ValueError("Модель Gemini не обрано. Будь ласка, оберіть модель у вкладці '⚙️ Settings'.")

    genai.configure(api_key=api_key.strip())

    directive = (
        "CRITICAL: You must return ONLY raw, valid JSON. Do not include any markdown formatting, "
        "do not use ```json blocks, and do not add any explanations, greetings, or conversational text "
        "before or after the JSON. Just the raw JSON object."
    )
    if directive not in system_prompt:
        system_prompt = system_prompt.strip() + " " + directive

    variables_json_str = json.dumps(variables_dict, ensure_ascii=False, indent=2)
    prompt_payload = (
        f"{system_prompt}\n\n"
        f"ГЛОБАЛЬНІ ЗМІННІ (словник):\n{variables_json_str}\n\n"
        f"ТЕКСТ ДОКУМЕНТА ДЛЯ АНАЛІЗУ:\n\"\"\"\n{document_text}\n\"\"\"\n\n"
        "ВАЖЛИВО: Повертай ТІЛЬКИ валідний JSON-об'єкт, де кожен ключ — це точний фрагмент тексту з документа, "
        "а значення — назва змінної з доступного словника. Без додаткових коментарів чи тексту навколо JSON."
    )

    model = genai.GenerativeModel(
        model_name=model_name.strip(),
        generation_config={"temperature": 0.1}
    )
    response = model.generate_content(prompt_payload, request_options={"timeout": 120})
    return parse_llm_json_response(response.text or "")


def query_local_llm_for_variables(
    base_url: str,
    model_name: str,
    system_prompt: str,
    variables_dict: Dict[str, str],
    document_text: str,
) -> Dict[str, str]:
    """
    Invokes Local LLM (Ollama) using the OpenAI SDK client.
    Sets a high timeout (600s) to handle heavy models (e.g. 27B) without timing out.
    """
    if not base_url.strip():
        raise ValueError("Local API Base URL не вказано. Будь ласка, введіть його у вкладці '⚙️ Settings'.")

    if not model_name.strip():
        raise ValueError("Модель Local LLM не обрано. Будь ласка, оберіть або введіть модель у вкладці '⚙️ Settings'.")

    # FIX: 600.0s timeout for heavy local models (Qwen 27B, Gemma 26B, etc.)
    client = OpenAI(
        base_url=base_url.strip(),
        api_key="ollama",
        timeout=600.0
    )

    directive = (
        "CRITICAL: You must return ONLY raw, valid JSON. Do not include any markdown formatting, "
        "do not use ```json blocks, and do not add any explanations, greetings, or conversational text "
        "before or after the JSON. Just the raw JSON object."
    )
    if directive not in system_prompt:
        system_prompt = system_prompt.strip() + " " + directive

    variables_json_str = json.dumps(variables_dict, ensure_ascii=False, indent=2)
    full_system_prompt = (
        f"{system_prompt}\n\n"
        f"ГЛОБАЛЬНІ ЗМІННІ (словник):\n{variables_json_str}\n\n"
        "ВАЖЛИВО: Повертай ТІЛЬКИ валідний JSON-об'єкт, де кожен ключ — це точний фрагмент тексту з документа, "
        "а значення — назва змінної з доступного словника. Без будь-яких додаткових коментарів чи тексту навколо JSON."
    )

    messages = [
        {"role": "system", "content": full_system_prompt},
        {"role": "user", "content": document_text},
    ]

    response = client.chat.completions.create(
        model=model_name.strip(),
        messages=messages,
        temperature=0.1,
    )

    raw_text = response.choices[0].message.content or ""
    return parse_llm_json_response(raw_text)


def query_variables(config: Dict[str, Any], document_text: str) -> Dict[str, str]:
    """Unified dispatcher for variable extraction based on configured provider."""
    provider = config.get("provider", DEFAULT_PROVIDER)
    system_prompt = config.get("system_prompt", DEFAULT_SYSTEM_PROMPT)
    variables_dict = config.get("variables", DEFAULT_VARIABLES)
    selected_model = config.get("selected_model", "")

    if provider == "Google Gemini":
        api_key = config.get("gemini_api_key", "")
        return query_gemini_for_variables(
            api_key=api_key,
            model_name=selected_model,
            system_prompt=system_prompt,
            variables_dict=variables_dict,
            document_text=document_text,
        )
    else:
        base_url = config.get("base_url", DEFAULT_LOCAL_BASE_URL)
        return query_local_llm_for_variables(
            base_url=base_url,
            model_name=selected_model,
            system_prompt=system_prompt,
            variables_dict=variables_dict,
            document_text=document_text,
        )


# ==========================================
# STREAMLIT UI APPLICATION
# ==========================================
def main():
    st.set_page_config(
        page_title="Legal Document Template Generator (Hybrid LLM)",
        page_icon="⚖️",
        layout="wide",
    )

    st.title("⚖️ Legal Document Template Generator")
    st.caption("Генератор шаблонів юридичних документів (.docx) з гібридною підтримкою Google Gemini та Local LLM (Ollama).")

    # ------------------------------------------
    # SESSION STATE INITIALIZATION (FIX FOR DISAPPEARING BUTTONS)
    # ------------------------------------------
    if "config" not in st.session_state:
        st.session_state.config = load_config()

    if "extracted_text" not in st.session_state:
        st.session_state.extracted_text = ""

    if "llm_parsed_json" not in st.session_state:
        st.session_state.llm_parsed_json = None

    if "document_name" not in st.session_state:
        st.session_state.document_name = None

    if "docx_bytes" not in st.session_state:
        st.session_state.docx_bytes = None

    if "editor_data" not in st.session_state:
        st.session_state.editor_data = pd.DataFrame(
            columns=["Keep", "Original Text", "Assigned Variable"]
        )

    if "generated_docx_buffer" not in st.session_state:
        st.session_state.generated_docx_buffer = None

    # Main Tabs
    tab_gen, tab_settings = st.tabs(["📄 Generation", "⚙️ Settings"])

    # ------------------------------------------
    # TAB 2: SETTINGS (HYBRID PROVIDER SUPPORT)
    # ------------------------------------------
    with tab_settings:
        st.subheader("⚙️ Конфігурація AI Провайдера та Моделі")
        st.markdown("Оберіть потрібного провайдера. Усі налаштування зберігаються локально в `config.json`.")

        current_provider = st.session_state.config.get("provider", DEFAULT_PROVIDER)
        current_gemini_key = st.session_state.config.get("gemini_api_key", "")
        current_base_url = st.session_state.config.get("base_url", DEFAULT_LOCAL_BASE_URL)
        current_selected_model = st.session_state.config.get("selected_model", DEFAULT_LOCAL_MODEL)
        current_prompt = st.session_state.config.get("system_prompt", DEFAULT_SYSTEM_PROMPT)
        current_variables = st.session_state.config.get("variables", DEFAULT_VARIABLES)

        # Provider Selector
        provider_options = ["Google Gemini", "Local LLM (Ollama)"]
        provider_idx = 0 if current_provider == "Google Gemini" else 1

        chosen_provider = st.radio(
            "1. AI Провайдер (AI Provider):",
            options=provider_options,
            index=provider_idx,
            horizontal=True,
            help="Виберіть між хмарним Google Gemini API або локальним сервером Ollama."
        )

        with st.form("settings_form"):
            # Provider-specific inputs
            if chosen_provider == "Google Gemini":
                st.markdown("#### Налаштування Google Gemini API")
                gemini_key_input = st.text_input(
                    "Gemini API Key",
                    value=current_gemini_key,
                    type="password",
                    help="Ваш персональний Google Gemini API Key."
                )

                # Fetch available Gemini models dynamically
                discovered_gemini = fetch_gemini_models(gemini_key_input.strip() or current_gemini_key)
                if discovered_gemini:
                    gemini_opts = list(discovered_gemini)
                    if current_selected_model not in gemini_opts:
                        gemini_opts.insert(0, current_selected_model)
                    model_idx = gemini_opts.index(current_selected_model) if current_selected_model in gemini_opts else 0

                    selected_model_input = st.selectbox(
                        "Робоча модель Gemini (отримано з ListModels)",
                        options=gemini_opts,
                        index=model_idx,
                        help="Динамічний список доступних моделей, що підтримують generateContent."
                    )
                else:
                    if gemini_key_input.strip():
                        st.caption("ℹ️ Натисніть 'Save Settings' для оновлення списку моделей через API.")
                    selected_model_input = st.text_input(
                        "Робоча модель Gemini",
                        value=current_selected_model if current_selected_model.startswith("models/") or "gemini" in current_selected_model else DEFAULT_GEMINI_MODEL,
                        help="Наприклад: models/gemini-3.5-flash-lite або models/gemini-2.5-flash"
                    )
                base_url_input = current_base_url  # preserve

            else:
                st.markdown("#### Налаштування Local LLM (Ollama / OpenAI API)")
                base_url_input = st.text_input(
                    "Local API Base URL",
                    value=current_base_url,
                    help="За замовчуванням: http://10.0.1.17:4333/v1 або http://brain.primocollect.ua/v1"
                )

                # Fetch available Local models dynamically
                discovered_local = fetch_local_models(base_url_input.strip())
                if discovered_local:
                    local_opts = list(discovered_local)
                    if current_selected_model not in local_opts:
                        local_opts.insert(0, current_selected_model)
                    model_idx = local_opts.index(current_selected_model) if current_selected_model in local_opts else 0

                    selected_model_input = st.selectbox(
                        "Робоча модель Local LLM (отримано з API)",
                        options=local_opts,
                        index=model_idx,
                        help="Список моделей, завантажених безпосередньо з вашого сервера Ollama."
                    )
                else:
                    st.caption("ℹ️ Сервер Local LLM недоступний або не відповідає на дану адресу. Введіть назву моделі вручну:")
                    selected_model_input = st.text_input(
                        "Робоча модель Local LLM (ручне введення)",
                        value=current_selected_model if "gemini" not in current_selected_model.lower() else DEFAULT_LOCAL_MODEL,
                        help="Наприклад: qwen3.8:27b або gemma4:26b-a4b-it-q8_0"
                    )
                gemini_key_input = current_gemini_key  # preserve

            st.markdown("#### Загальні налаштування юридичного аналізу")
            prompt_input = st.text_area(
                "Системний промпт (System Prompt)",
                value=current_prompt,
                height=130,
                help="Інструкція для моделі, що визначає правила точного виділення змінних."
            )

            variables_json_str = json.dumps(current_variables, ensure_ascii=False, indent=2)
            variables_input = st.text_area(
                "Словник глобальних змінних (JSON)",
                value=variables_json_str,
                height=220,
                help="JSON-словник змінних. Ключі відображатимуться у випадаючому списку вибору."
            )

            submitted = st.form_submit_button("💾 Save Settings", use_container_width=True)

            if submitted:
                try:
                    parsed_vars = json.loads(variables_input)
                    if not isinstance(parsed_vars, dict):
                        st.error("Помилка: Словник змінних має бути валідним JSON-об'єктом.")
                    else:
                        new_config = {
                            "provider": chosen_provider,
                            "gemini_api_key": gemini_key_input.strip(),
                            "base_url": base_url_input.strip(),
                            "selected_model": selected_model_input.strip(),
                            "system_prompt": prompt_input.strip(),
                            "variables": parsed_vars,
                        }
                        if save_config(new_config):
                            st.session_state.config = new_config
                            st.success(f"✅ Налаштування збережено! Провайдер: `{chosen_provider}` | Модель: `{selected_model_input.strip()}`")
                            st.rerun()
                except json.JSONDecodeError as err:
                    st.error(f"❌ Помилка валідації JSON у словнику змінних: {err}")

    # ------------------------------------------
    # TAB 1: GENERATION (PERSISTENT UI WORKFLOW)
    # ------------------------------------------
    with tab_gen:
        st.subheader("📄 Генерація шаблону DOCX")

        active_provider = st.session_state.config.get("provider", DEFAULT_PROVIDER)
        active_model = st.session_state.config.get("selected_model", "")
        var_dict = st.session_state.config.get("variables", DEFAULT_VARIABLES)
        var_options = get_variable_names(var_dict)

        st.caption(f"Поточний провайдер: **{active_provider}** | Модель: `{active_model}`")

        # Step 1: File Uploader
        uploaded_file = st.file_uploader(
            "Завантажте документ у форматі .docx",
            type=["docx"],
            help="Оберіть файл .docx для автоматичного аналізу та перетворення на шаблон."
        )

        # Handle file upload state persistence
        if uploaded_file is not None:
            if st.session_state.document_name != uploaded_file.name:
                st.session_state.document_name = uploaded_file.name
                file_bytes = uploaded_file.getvalue()
                st.session_state.docx_bytes = file_bytes
                try:
                    doc = Document(io.BytesIO(file_bytes))
                    st.session_state.extracted_text = extract_text_from_docx(doc)
                except Exception as e:
                    st.error(f"Помилка читання .docx файлу: {e}")
                    st.session_state.extracted_text = ""
                # Reset analysis results when a brand-new file is provided
                st.session_state.llm_parsed_json = None
                st.session_state.editor_data = pd.DataFrame(
                    columns=["Keep", "Original Text", "Assigned Variable"]
                )
                st.session_state.generated_docx_buffer = None
        else:
            if st.session_state.document_name is not None:
                st.session_state.document_name = None
                st.session_state.docx_bytes = None
                st.session_state.extracted_text = ""
                st.session_state.llm_parsed_json = None
                st.session_state.editor_data = pd.DataFrame(
                    columns=["Keep", "Original Text", "Assigned Variable"]
                )
                st.session_state.generated_docx_buffer = None

        # When a file is currently uploaded and parsed
        if st.session_state.docx_bytes is not None:
            with st.expander(f"🔍 Попередній перегляд вилученого тексту ({st.session_state.document_name})", expanded=False):
                st.text_area("Вилучений текст (параграфи + таблиці)", st.session_state.extracted_text, height=180, disabled=True)

            col_run, col_clear = st.columns([3, 1])
            with col_run:
                run_analysis = st.button(
                    f"🤖 Аналізувати документ за допомогою {active_provider}",
                    use_container_width=True,
                    type="primary"
                )
            with col_clear:
                if st.button("🔄 Очистити аналіз", use_container_width=True):
                    st.session_state.llm_parsed_json = None
                    st.session_state.editor_data = pd.DataFrame(
                        columns=["Keep", "Original Text", "Assigned Variable"]
                    )
                    st.session_state.generated_docx_buffer = None
                    st.rerun()

            # Execute Analysis
            if run_analysis:
                if not st.session_state.extracted_text.strip():
                    st.error("Помилка: Не вдалося вилучити текст із завантаженого .docx документа.")
                else:
                    with st.spinner(f"Запит до `{active_model}` ({active_provider})... Аналізуємо документ та шукаємо змінні..."):
                        try:
                            llm_mapping = query_variables(
                                config=st.session_state.config,
                                document_text=st.session_state.extracted_text
                            )

                            # Save to session_state
                            st.session_state.llm_parsed_json = llm_mapping

                            rows = []
                            for snippet, assigned_var in llm_mapping.items():
                                valid_var = assigned_var if assigned_var in var_options else (var_options[0] if var_options else "")
                                rows.append({
                                    "Keep": True,
                                    "Original Text": str(snippet),
                                    "Assigned Variable": valid_var
                                })

                            if not rows:
                                st.info("ℹ️ Модель не знайшла збігів зі словником змінних у цьому документі.")
                                st.session_state.editor_data = pd.DataFrame(
                                    columns=["Keep", "Original Text", "Assigned Variable"]
                                )
                            else:
                                st.session_state.editor_data = pd.DataFrame(rows)
                                st.success(f"✅ Знайдено {len(rows)} потенційних змінних!")

                            st.session_state.generated_docx_buffer = None
                            st.rerun()

                        except Exception as ex:
                            err_msg = str(ex)
                            if "connection" in err_msg.lower() or "connect" in err_msg.lower():
                                st.error(
                                    f"❌ Помилка з'єднання з сервером LLM. "
                                    f"Перевірте підключення до мережі та параметри у вкладці Settings.\n\nДеталі: {err_msg}"
                                )
                            elif "timed out" in err_msg.lower() or "timeout" in err_msg.lower():
                                st.error(
                                    f"⏱️ Час очікування відповіді від моделі `{active_model}` вичерпано. "
                                    f"Для великих моделей (27B) встановлено таймаут 600 секунд, але сервер не відповів вчасно."
                                )
                            else:
                                st.error(f"❌ Помилка аналізу через {active_provider}: {ex}")

        # Step 2 & 3: Human-in-the-Loop & Template Generation (RENDER ONLY IF llm_parsed_json IS NOT NONE)
        if st.session_state.llm_parsed_json is not None:
            st.markdown("---")
            st.markdown("### 🧑‍⚖️ Human-in-the-Loop: Перевірка та коригування змінних")
            st.caption("Позначте 'Keep' для заміни фрагмента. За потреби змініть назву змінної у випадаючому списку.")

            if not var_options:
                st.warning("Увага: Словник змінних порожній! Додайте змінні у вкладці 'Settings'.")

            column_config = {
                "Keep": st.column_config.CheckboxColumn(
                    "Keep",
                    help="Позначте для заміни цього фрагмента на шаблонну змінну",
                    default=True,
                ),
                "Original Text": st.column_config.TextColumn(
                    "Original Text",
                    help="Оригінальний фрагмент тексту з документа для заміни",
                    required=True,
                ),
                "Assigned Variable": st.column_config.SelectboxColumn(
                    "Assigned Variable",
                    help="Призначена змінна зі списку доступних у конфігурації",
                    width="medium",
                    options=var_options,
                    required=True,
                ),
            }

            edited_df = st.data_editor(
                st.session_state.editor_data,
                column_config=column_config,
                disabled=False,
                num_rows="dynamic",
                use_container_width=True,
                key="data_editor"
            )

            # Step 4: Template Generation
            st.markdown("---")
            st.subheader("⚡ Генерація та завантаження шаблону DOCX")

            col_gen, col_download = st.columns([2, 2])

            with col_gen:
                if st.button("🚀 Generate Template", use_container_width=True, type="primary"):
                    if not st.session_state.docx_bytes:
                        st.warning("Будь ласка, завантажте .docx файл для генерації шаблону.")
                    elif edited_df.empty:
                        st.warning("Таблиця змінних порожня. Додайте змінні або виконайте аналіз документа.")
                    else:
                        items_to_replace = []
                        for _, row in edited_df.iterrows():
                            keep = row.get("Keep", False)
                            orig_text = str(row.get("Original Text", "")).strip()
                            var_name = str(row.get("Assigned Variable", "")).strip()
                            if keep and orig_text and var_name:
                                items_to_replace.append((orig_text, var_name))

                        if not items_to_replace:
                            st.warning("Жодного рядка не позначено для заміни (Keep = True).")
                        else:
                            with st.spinner("Застосовуємо заміни та жовте підсвічування до .docx..."):
                                fresh_doc = Document(io.BytesIO(st.session_state.docx_bytes))
                                mod_doc, count = generate_template_docx(fresh_doc, items_to_replace)

                                out_buffer = io.BytesIO()
                                mod_doc.save(out_buffer)
                                out_buffer.seek(0)
                                st.session_state.generated_docx_buffer = out_buffer.getvalue()

                                st.success(
                                    f"🎉 Шаблон успішно згенеровано! Замінено {count} фрагментів "
                                    f"із жовтим підсвічуванням (WD_COLOR_INDEX.YELLOW)."
                                )
                                st.rerun()

            with col_download:
                if st.session_state.generated_docx_buffer:
                    st.download_button(
                        label="📥 Завантажити Template.docx",
                        data=st.session_state.generated_docx_buffer,
                        file_name="Template.docx",
                        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                        use_container_width=True
                    )
                else:
                    st.button("📥 Завантажити Template.docx", disabled=True, use_container_width=True)

        elif st.session_state.docx_bytes is not None:
            st.info("👆 Натисніть кнопку '🤖 Аналізувати документ...', щоб вилучити змінні за допомогою обраної моделі.")


if __name__ == "__main__":
    main()
