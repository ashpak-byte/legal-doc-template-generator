import io
from pathlib import Path
import docx
from docx import Document
from docx.enum.text import WD_COLOR_INDEX
import app

def test_pipeline():
    sample_file = Path("sample_contract.docx")
    assert sample_file.exists(), "sample_contract.docx must exist for pipeline verification"

    doc = Document(str(sample_file))

    # 1. Text extraction
    plain_text = app.extract_text_from_docx(doc)
    print("=== Extracted Plain Text Sample ===")
    print(plain_text[:200])
    assert "4820-2024" in plain_text
    assert "Коваленко Петро Васильович" in plain_text

    # 2. Simulated LLM extraction results based on sample_contract.docx
    mock_llm_mapping = [
        ("4820-2024", "pt_agreement_number"),
        ("Коваленко Петро Васильович", "pt_last_name"),
        ("3123456789", "pt_pin"),
        ("м. Київ, вул. Хрещатик, буд. 22, кв. 10", "pt_address"),
        ("+380501234567", "pt_phone"),
        ("45000", "pt_capital"),
        ("52300", "pt_total_debt"),
        ("3028", "pt_court_fee"),
    ]

    # 3. Generate template
    mod_doc, count = app.generate_template_docx(doc, mock_llm_mapping)
    print(f"\nGenerated DOCX: replaced_count={count}")
    assert count >= len(mock_llm_mapping), f"Expected at least {len(mock_llm_mapping)} replacements, got {count}"

    # 4. Save to in-memory buffer
    out_buf = io.BytesIO()
    mod_doc.save(out_buf)
    out_buf.seek(0)
    assert len(out_buf.getvalue()) > 0

    # 5. Reload and verify yellow highlights
    reloaded = Document(out_buf)
    all_text = app.extract_text_from_docx(reloaded)
    assert "$$$pt_agreement_number$$$" in all_text
    assert "$$$pt_total_debt$$$" in all_text

    # Count yellow highlighted runs across all paragraphs and tables
    yellow_runs = 0
    for p in reloaded.paragraphs:
        for r in p.runs:
            if r.font.highlight_color == WD_COLOR_INDEX.YELLOW:
                yellow_runs += 1

    for table in reloaded.tables:
        for row in table.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    for r in p.runs:
                        if r.font.highlight_color == WD_COLOR_INDEX.YELLOW:
                            yellow_runs += 1

    print(f"Yellow highlighted runs found in document: {yellow_runs}")
    assert yellow_runs >= len(mock_llm_mapping)

    print("\n[OK] End-to-end DOCX pipeline test passed successfully!")

if __name__ == "__main__":
    test_pipeline()
